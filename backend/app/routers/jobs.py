import uuid
from datetime import datetime, timezone
from fastapi import APIRouter, Depends, HTTPException, Response, status
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from app.database import get_db
from app.dependencies.auth import get_current_user
from app.schemas.job import (
    JobCreate,
    JobUpdate,
    JobOut,
    JobCardOut,
    JobEventOut,
    FollowUpUpdate,
    JobDetailOut,
    JobSearchParams,
    BulkStatusUpdate,
    BulkDelete,
    JobArtifactCreate,
    JobArtifactOut,
)
from app.models.job import Job, JobStatus, ArtifactKind
from app.services import jobs as jobs_service
from app.services import workspace as workspace_service
from app.services import company as company_service
from app.services import metrics as metrics_service
from pydantic import BaseModel

router = APIRouter(prefix="/api/v1/jobs", tags=["jobs"])


async def get_workspace_id(user: dict, db: AsyncSession) -> uuid.UUID:
    ws = await workspace_service.get_workspace_for_user(db, uuid.UUID(user["user_id"]))
    if not ws:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Workspace not found")
    return ws.id


@router.get("/", response_model=list[JobOut])
async def list_jobs(
    board_id: uuid.UUID | None = None,
    search: str | None = None,
    status: str | None = None,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    try:
        status_enum = JobStatus(status) if status else None
    except ValueError:
        # `status` is shadowed by the query param here, so use the literal code.
        raise HTTPException(status_code=400, detail=f"Unknown status '{status}'")
    return await jobs_service.list_jobs(db, ws_id, board_id, search, status_enum)


def _duplicate_409(exc: jobs_service.DuplicateJobError) -> HTTPException:
    return HTTPException(
        status_code=status.HTTP_409_CONFLICT,
        detail={"message": str(exc), "existing_job_id": str(exc.existing_id)},
    )


@router.get("/cards", response_model=list[JobCardOut])
async def list_job_cards(
    response: Response,
    board_id: uuid.UUID | None = None,
    search: str | None = None,
    status: str | None = None,
    source: str | None = None,
    min_fit: int | None = None,
    location: str | None = None,
    limit: int = 500,
    offset: int = 0,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Slim job list for boards and tables. `status` may be a comma-separated list. The total
    number of matches (ignoring limit/offset) is returned in the X-Total-Count header."""
    ws_id = await get_workspace_id(user, db)
    statuses = None
    if status:
        statuses = [s.strip() for s in status.split(",") if s.strip()]
        try:
            statuses = [JobStatus(s).value for s in statuses]
        except ValueError:
            raise HTTPException(status_code=400, detail=f"Unknown status in '{status}'")
    rows, total = await jobs_service.list_cards(
        db, ws_id,
        board_id=board_id, search=search, statuses=statuses, source=source,
        min_fit=min_fit, location=location,
        limit=max(1, min(limit, 1000)), offset=max(0, offset),
    )
    response.headers["X-Total-Count"] = str(total)
    return rows


@router.get("/stats")
async def job_stats(
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    ws = await workspace_service.get_workspace_for_user(db, uuid.UUID(user["user_id"]))
    return await metrics_service.job_stats(db, ws_id, datetime.now(timezone.utc), ws.timezone if ws else None)


@router.get("/events", response_model=list[JobEventOut])
async def list_events(
    limit: int = 100,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    return await jobs_service.list_job_events(db, ws_id, limit=max(1, min(limit, 500)))


@router.get("/export")
async def export_jobs(
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    import io
    from datetime import datetime as _dt

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ws_id = await get_workspace_id(user, db)
    jobs = await jobs_service.list_jobs(db, ws_id)

    headers = [
        "Company", "Title", "Location", "Status", "Date Applied",
        "Company Domain", "Job URL", "Resume Score", "Notes", "Description",
        "Insights", "Base Fit Score", "Source", "Applied At", "Follow-up", "Created At", "Updated At",
    ]

    def _flatten(value):
        if value is None:
            return ""
        if isinstance(value, dict):
            return "\n".join(f"{k}: {v}" for k, v in value.items())
        if isinstance(value, list):
            return ", ".join(str(x) for x in value)
        return str(value)

    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"

    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(bold=True, color="FFFFFF")

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for job in jobs:
        ws.append([
            job.company or "",
            job.title or "",
            job.location or "",
            job.status or "",
            str(job.date_applied) if job.date_applied else "",
            job.company_domain or "",
            job.url or "",
            job.resume_score if job.resume_score is not None else "",
            job.notes or "",
            job.description or "",
            _flatten(job.insights),
            _flatten(job.base_fit_score),
            job.source or "",
            str(job.applied_at) if job.applied_at else "",
            str(job.follow_up_at) if job.follow_up_at else "",
            str(job.created_at) if job.created_at else "",
            str(job.updated_at) if job.updated_at else "",
        ])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for i, w in enumerate([18, 22, 16, 12, 13, 20, 30, 13, 30, 50, 40, 40, 14, 22, 14, 22, 22], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"jobforge_jobs_{_dt.now():%Y-%m-%d}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{job_id}", response_model=JobDetailOut)
async def get_job(
    job_id: uuid.UUID,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    try:
        return await jobs_service.get_job_detail(db, ws_id, job_id)
    except jobs_service.NotFoundError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e))


@router.get("/{job_id}/company-info")
async def get_company_info(
    job_id: uuid.UUID,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    result = await db.execute(
        select(Job).where(Job.id == job_id, Job.workspace_id == ws_id)
    )
    job = result.scalar_one_or_none()
    if not job:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Job not found")

    domain = company_service.resolve_domain(
        company=job.company, domain=job.company_domain, url=job.url
    )
    website = f"https://{domain}"

    description = await company_service.fetch_website_description(domain)
    if description:
        return {
            "company": job.company,
            "domain": domain,
            "website": website,
            "description": description,
            "url": None,
            "source": "website",
        }

    wiki = await company_service.wikipedia_lookup(job.company, domain)
    if wiki:
        return {
            "company": job.company,
            "domain": domain,
            "website": website,
            "description": wiki["description"],
            "url": wiki.get("url"),
            "source": "wikipedia",
        }

    return {
        "company": job.company,
        "domain": domain,
        "website": website,
        "description": None,
        "url": None,
        "source": None,
    }


@router.post("/", response_model=JobOut)
async def create_job(
    data: JobCreate,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    try:
        job = await jobs_service.create_job(
            db,
            workspace_id=ws_id,
            board_id=data.board_id,
            company=data.company,
            title=data.title,
            description=data.description or "",
            url=data.url,
            notes=data.notes,
            location=data.location,
            status=data.status,
            date_applied=str(data.date_applied) if data.date_applied else None,
            resume_score=data.resume_score,
            company_domain=data.company_domain,
        )
    except jobs_service.NotFoundError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e))
    except jobs_service.DuplicateJobError as e:
        raise _duplicate_409(e)
    await db.refresh(job)
    return job


@router.put("/{job_id}")
async def update_job(
    job_id: uuid.UUID,
    data: JobUpdate,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    updates = {k: v for k, v in data.model_dump().items() if v is not None}
    try:
        await jobs_service.update_job(db, ws_id, job_id, **updates)
    except jobs_service.NotFoundError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e))
    except jobs_service.DuplicateJobError as e:
        raise _duplicate_409(e)
    return {"ok": True}


@router.put("/{job_id}/follow-up")
async def set_follow_up(
    job_id: uuid.UUID,
    data: FollowUpUpdate,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Set (or clear, with null) the date to chase this application."""
    ws_id = await get_workspace_id(user, db)
    try:
        job = await jobs_service.set_follow_up(db, ws_id, job_id, data.follow_up_at)
    except jobs_service.NotFoundError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e))
    return {"ok": True, "follow_up_at": job.follow_up_at}


@router.get("/{job_id}/events", response_model=list[JobEventOut])
async def list_job_events(
    job_id: uuid.UUID,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    if await jobs_service.get_job(db, ws_id, job_id) is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Job not found")
    return await jobs_service.list_job_events(db, ws_id, job_id=job_id, limit=200)


@router.delete("/{job_id}")
async def delete_job(
    job_id: uuid.UUID,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    await jobs_service.delete_job(db, ws_id, job_id)
    return {"ok": True}


@router.post("/bulk-status")
async def bulk_update_status(
    data: BulkStatusUpdate,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    await jobs_service.bulk_update_status(db, ws_id, data.ids, data.status)
    return {"ok": True}


@router.post("/bulk-delete")
async def bulk_delete_jobs(
    data: BulkDelete,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    count = await jobs_service.bulk_delete_jobs(db, ws_id, data.ids)
    return {"ok": True, "count": count}


@router.get("/artifacts/{job_id}", response_model=list[JobArtifactOut])
async def list_artifacts(
    job_id: uuid.UUID,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    try:
        detail = await jobs_service.get_job_detail(db, ws_id, job_id)
    except jobs_service.NotFoundError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e))
    return detail["artifacts"]


@router.post("/artifacts", response_model=JobArtifactOut)
async def create_artifact(
    data: JobArtifactCreate,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    try:
        artifact = await jobs_service.create_artifact(
            db,
            workspace_id=ws_id,
            job_id=data.job_id,
            kind=ArtifactKind(data.kind.value),
            filename=data.filename,
            latex_source=data.latex_source,
        )
    except jobs_service.NotFoundError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e))
    await db.refresh(artifact)
    return artifact


@router.delete("/artifacts/{artifact_id}")
async def delete_artifact(
    artifact_id: uuid.UUID,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    await jobs_service.delete_artifact(db, ws_id, artifact_id)
    return {"ok": True}


class ArtifactUpdateRequest(BaseModel):
    latex_source: str


@router.patch("/artifacts/{artifact_id}")
async def update_artifact(
    artifact_id: uuid.UUID,
    data: ArtifactUpdateRequest,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    ws_id = await get_workspace_id(user, db)
    try:
        art = await jobs_service.update_artifact_source(db, ws_id, artifact_id, data.latex_source)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    if not art:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Artifact not found")
    return {"ok": True}
